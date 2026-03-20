#!/usr/bin/env python3
"""
Generator pentru fisierul Excel 'NecesarAprovizionare.xlsx'
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# -- Stiluri -----------------------------------------------------------------------
HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
BTN_FONT = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
BTN_ALIGN = Alignment(horizontal="center", vertical="center")
SUBTITLE_FONT = Font(name="Calibri", size=10, color="666666")
SECTION_FONT = Font(name="Calibri", bold=True, size=13, color="2F5496")
CAT_HEADER_FILL = PatternFill(start_color="F57C00", end_color="F57C00", fill_type="solid")
CAT_HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")


def style_header_row(ws, row, col_start, col_end, fill=None, font=None):
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = font or HEADER_FONT
        cell.fill = fill or HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


def fill_btn(ws, cell_ref, text, color):
    cell = ws[cell_ref]
    cell.value = text
    cell.font = BTN_FONT
    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    cell.alignment = BTN_ALIGN
    cell.border = THIN_BORDER


def create_menu_sheet(wb):
    ws = wb.active
    ws.title = "Meniu"
    ws.sheet_properties.tabColor = "2F5496"
    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 3

    # Titlu
    ws.merge_cells("B2:C2")
    ws["B2"].value = "NECESAR APROVIZIONARE - MAGAZIN"
    ws["B2"].font = Font(name="Calibri", bold=True, size=18, color="2F5496")
    ws["B2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("B3:C3")
    ws["B3"].value = "Sistem de calcul necesar aprovizionare"
    ws["B3"].font = SUBTITLE_FONT
    ws["B3"].alignment = Alignment(horizontal="center")

    # PASUL 1
    ws.merge_cells("B5:C5")
    ws["B5"].value = "PASUL 1: IMPORT DATE"
    ws["B5"].font = SECTION_FONT

    # Buton Import Stoc Magazin - B7:B8
    ws.merge_cells("B7:B8")
    fill_btn(ws, "B7", "IMPORT STOC MAGAZIN", "2E7D32")
    ws.merge_cells("C7:C8")
    ws["C7"].value = "Importa stocul magazinului. Format: CodProdus | Stoc"
    ws["C7"].font = Font(name="Calibri", size=10, color="333333")
    ws["C7"].alignment = Alignment(vertical="center", wrap_text=True)

    # Buton Import Vanzari - B10:B11
    ws.merge_cells("B10:B11")
    fill_btn(ws, "B10", "IMPORT VANZARI MAGAZIN", "1565C0")
    ws.merge_cells("C10:C11")
    ws["C10"].value = "Importa vanzarile magazinului. Format: CodProdus | Cantitate"
    ws["C10"].font = Font(name="Calibri", size=10, color="333333")
    ws["C10"].alignment = Alignment(vertical="center", wrap_text=True)

    # PASUL 2
    ws.merge_cells("B13:C13")
    ws["B13"].value = "PASUL 2: GENERARE NECESAR"
    ws["B13"].font = SECTION_FONT

    # Buton Generare Necesar - B14:B15
    ws.merge_cells("B14:B15")
    fill_btn(ws, "B14", "GENERARE NECESAR", "6A1B9A")
    ws.merge_cells("C14:C15")
    ws["C14"].value = "Calculeaza necesarul cu compatibilitati si rotunjiri."
    ws["C14"].font = Font(name="Calibri", size=10, color="333333")
    ws["C14"].alignment = Alignment(vertical="center", wrap_text=True)

    # ACTIUNI
    ws.merge_cells("B17:C17")
    ws["B17"].value = "ACTIUNI"
    ws["B17"].font = SECTION_FONT

    # Buton Sterge - B18:B19
    ws.merge_cells("B18:B19")
    fill_btn(ws, "B18", "STERGE TOATE DATELE", "C62828")
    ws.merge_cells("C18:C19")
    ws["C18"].value = "Sterge datele importate si necesarul generat."
    ws["C18"].font = Font(name="Calibri", size=10, color="333333")
    ws["C18"].alignment = Alignment(vertical="center", wrap_text=True)

    # STATUS - VBA expects: C20=stoc, C21=vanzari, C22=ultima actualizare, C23=necesar
    ws["B20"].value = "Stoc Magazin:"
    ws["B20"].font = Font(name="Calibri", bold=True, size=11)
    ws["C20"].value = "Neincarcat"
    ws["C20"].font = Font(name="Calibri", size=11, color="C62828")

    ws["B21"].value = "Vanzari Magazin:"
    ws["B21"].font = Font(name="Calibri", bold=True, size=11)
    ws["C21"].value = "Neincarcat"
    ws["C21"].font = Font(name="Calibri", size=11, color="C62828")

    ws["B22"].value = "Ultima actualizare:"
    ws["B22"].font = Font(name="Calibri", bold=True, size=11)
    ws["C22"].value = "-"
    ws["C22"].font = Font(name="Calibri", size=11, color="666666")

    ws["B23"].value = "Necesar:"
    ws["B23"].font = Font(name="Calibri", bold=True, size=11)
    ws["C23"].value = "Negenerat"
    ws["C23"].font = Font(name="Calibri", size=11, color="C62828")

    # Instructiuni
    ws.merge_cells("B25:C25")
    ws["B25"].value = "INSTRUCTIUNI"
    ws["B25"].font = SECTION_FONT

    instructions = [
        "1. Importati stocul si vanzarile folosind butoanele de mai sus.",
        "2. Fisierele trebuie sa respecte sabloanele din folderul 'sabloane'.",
        "3. Configurati Compatibilitati si ReguliRotunjire daca e necesar.",
        "4. Setati zilele de vanzare si necesar mai jos.",
        "5. Apasati GENERARE NECESAR pentru a calcula necesarul.",
    ]
    for i, text in enumerate(instructions):
        r = 27 + i
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"].value = text
        ws[f"B{r}"].font = Font(name="Calibri", size=10, color="444444")

    # -- SETARI CALCUL NECESAR (VBA citeste C33 si C34) --
    settings_fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
    settings_border = THIN_BORDER

    ws["B33"].value = "Zile vanzare:"
    ws["B33"].font = Font(name="Calibri", bold=True, size=11)
    ws["C33"].value = 30
    ws["C33"].font = Font(name="Calibri", bold=True, size=14, color="2F5496")
    ws["C33"].fill = settings_fill
    ws["C33"].border = settings_border
    ws["C33"].alignment = Alignment(horizontal="center")

    ws["B34"].value = "Zile necesar:"
    ws["B34"].font = Font(name="Calibri", bold=True, size=11)
    ws["C34"].value = 7
    ws["C34"].font = Font(name="Calibri", bold=True, size=14, color="2F5496")
    ws["C34"].fill = settings_fill
    ws["C34"].border = settings_border
    ws["C34"].alignment = Alignment(horizontal="center")

    ws.merge_cells("B32:C32")
    ws["B32"].value = "SETARI CALCUL"
    ws["B32"].font = SECTION_FONT

    # Explicatie formula
    ws.merge_cells("B36:C36")
    ws["B36"].value = "Formula: Necesar = CEILING((Vanzare / ZileVanz) * ZileNec) - Stoc"
    ws["B36"].font = Font(name="Calibri", italic=True, size=9, color="888888")

    return ws


def create_stoc_magazin_sheet(wb):
    ws = wb.create_sheet("StocMagazin")
    ws.sheet_properties.tabColor = "2E7D32"
    for i, (h, w) in enumerate([("Cod Produs", 25), ("Stoc Magazin", 20)], 1):
        ws.cell(row=1, column=i, value=h)
        ws.column_dimensions[chr(64 + i)].width = w
    style_header_row(ws, 1, 1, 2)
    ws.auto_filter.ref = "A1:B1"
    return ws


def create_vanzari_magazin_sheet(wb):
    ws = wb.create_sheet("VanzariMagazin")
    ws.sheet_properties.tabColor = "1565C0"
    for i, (h, w) in enumerate([("Cod Produs", 25), ("Cantitate Vanzare", 22)], 1):
        ws.cell(row=1, column=i, value=h)
        ws.column_dimensions[chr(64 + i)].width = w
    style_header_row(ws, 1, 1, 2)
    ws.auto_filter.ref = "A1:B1"
    return ws


def create_compatibilitati_sheet(wb):
    """4 coloane: Cod Original | Denumire Original | Cod Compatibil | Denumire Compatibil"""
    ws = wb.create_sheet("Compatibilitati")
    ws.sheet_properties.tabColor = "F57C00"

    headers = [
        ("Cod Produs Original", 22),
        ("Denumire Produs Original", 50),
        ("Cod Produs Compatibil", 22),
        ("Denumire Produs Compatibil", 50),
    ]
    for i, (h, w) in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
        ws.column_dimensions[chr(64 + i)].width = w

    style_header_row(ws, 1, 1, 4)
    ws.auto_filter.ref = "A1:D1"
    return ws


def create_reguli_rotunjire_sheet(wb):
    """
    Sectiunea 1 (rand 1-5): Categorii de rotunjire
    Sectiunea 2 (rand 7+): Lista produse cu categorie asignata
    """
    ws = wb.create_sheet("ReguliRotunjire")
    ws.sheet_properties.tabColor = "7B1FA2"

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 18

    # Sectiunea 1: Categorii
    ws.cell(row=1, column=1, value="Categorie")
    ws.cell(row=1, column=2, value="Descriere")
    style_header_row(ws, 1, 1, 2, fill=CAT_HEADER_FILL, font=CAT_HEADER_FONT)

    categories = [
        (1, "Rotunjire la multiplu de 10. Valori sub 5 devin 10. Exemplu: 3->10, 12->10, 15->20, 27->30"),
        (2, "Rotunjire la numar par (multiplu de 2). Exemplu: 1->2, 3->4, 5->6, 7->8"),
        (3, "Rotunjire la multiplu de 5. Valori 1-2 devin 0 (se elimina). Exemplu: 2->0, 3->5, 7->5, 8->10"),
        (4, "Rotunjire la multiplu de 5 cu minim 5 daca stoc magazin < 3. Exemplu: 2->5(daca stoc<3), 8->10"),
    ]
    cat_font = Font(name="Calibri", size=10)
    cat_num_font = Font(name="Calibri", bold=True, size=12)
    for r, (cat_num, desc) in enumerate(categories, 2):
        ws.cell(row=r, column=1, value=cat_num).font = cat_num_font
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=1).border = THIN_BORDER
        ws.cell(row=r, column=2, value=desc).font = cat_font
        ws.cell(row=r, column=2).alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row=r, column=2).border = THIN_BORDER
        ws.row_dimensions[r].height = 30

    # Sectiunea 2: Produse cu categorie asignata (de la randul 7)
    ws.cell(row=7, column=1, value="Cod")
    ws.cell(row=7, column=2, value="Denumire")
    ws.cell(row=7, column=3, value="Categorie")
    style_header_row(ws, 7, 1, 3)

    return ws


def create_necesar_sheet(wb):
    ws = wb.create_sheet("Necesar")
    ws.sheet_properties.tabColor = "6A1B9A"

    headers = [
        ("Cod Produs", 25),
        ("Vanzare", 15),
        ("Medie/Zi", 12),
        ("Stoc Magazin", 15),
        ("Necesar", 15),
        ("Observatii", 55),
    ]
    for i, (h, w) in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
        ws.column_dimensions[chr(64 + i)].width = w

    style_header_row(ws, 1, 1, len(headers))
    ws.auto_filter.ref = f"A1:{chr(64 + len(headers))}1"
    return ws


def create_log_sheet(wb):
    ws = wb.create_sheet("Log")
    ws.sheet_properties.tabColor = "757575"
    for i, (h, w) in enumerate([
        ("Data/Ora", 22), ("Operatiune", 25), ("Detalii", 50), ("Status", 15)
    ], 1):
        ws.cell(row=1, column=i, value=h)
        ws.column_dimensions[chr(64 + i)].width = w
    style_header_row(ws, 1, 1, 4)
    return ws


def main():
    wb = openpyxl.Workbook()

    create_menu_sheet(wb)
    create_stoc_magazin_sheet(wb)
    create_vanzari_magazin_sheet(wb)
    create_compatibilitati_sheet(wb)
    create_reguli_rotunjire_sheet(wb)
    create_necesar_sheet(wb)
    create_log_sheet(wb)

    output_path = "NecesarAprovizionare.xlsx"
    wb.save(output_path)
    print(f"Fisierul '{output_path}' a fost generat cu succes!")
    print()
    print("SETUP (3 pasi):")
    print("  1. Deschideti in Excel -> Save As -> .xlsm")
    print("  2. Alt+F11 -> Import File -> vba_modules/ModNecesar.bas")
    print("  3. Alt+F8 -> ConfigureazaAplicatia -> Run -> Save")


if __name__ == "__main__":
    main()
